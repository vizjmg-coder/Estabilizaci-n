import json
import openpyxl

# 1. Read Excel raw sheet
wb = openpyxl.load_workbook('Data/C_Estabilizacion_633km_subregion_tabla.xlsx', data_only=True)
print("Sheet names:", wb.sheetnames)
ws = wb['SUBREGIONES']

excel_circuits = []
for row in ws.iter_rows(min_row=1, max_row=40, values_only=True):
    if row[0] is not None:
        excel_circuits.append(row)

for r in excel_circuits[:5]:
    print(r)

# 2. Check all municipalities mentioned in Excel
with open('Data/data_normalized.json', 'r', encoding='utf-8') as f:
    data = json.load(f)

excel_muns = set()
for c in data['circuits']:
    for m in c['municipios']:
        excel_muns.add(m.strip().upper())

print("\n--- Excel Municipalities ---")
print("Total unique in Excel:", len(excel_muns))
print("Sorted list:", sorted(list(excel_muns)))

print("\n--- Checking specific ones ---")
print("MEDELLIN in Excel?:", 'MEDELLÍN' in excel_muns or 'MEDELLIN' in excel_muns)
print("GIRALDO in Excel?:", 'GIRALDO' in excel_muns)

# 3. Check Municipios_opt.geojson
with open('Data/Municipios_opt.geojson', 'r', encoding='utf-8') as f:
    mun_geo = json.load(f)

print("\n--- Checking in Municipios_opt.geojson ---")
for f in mun_geo['features']:
    p = f['properties']
    name = p.get('name', '').upper()
    if 'MEDELL' in name or 'GIRALDO' in name:
        print(f"Name: {p.get('name')} | Subreg: {p.get('subregion')} | Circuits count: {p.get('circuits_count')} | Km: {p.get('km_intervencion')}")
